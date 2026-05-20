# tests/test_security_controls.py
# Focused DevSecOps/security regression tests

from io import BytesIO
from types import SimpleNamespace
from bson import ObjectId
from docx import Document

import app as cybermailguard_app
from app import create_app
from src.auth.twofa import generate_secret, verify_token
from src.pattern_engine.ml_classifier import train_model
from src.dashboard import routes as dashboard_routes


def make_app(csrf_enabled=True):
    return create_app({
        'TESTING': True,
        'SECRET_KEY': 'test-secret',
        'MONGO_URI': 'mongodb://localhost:27017/cybermailguard_test',
        'WTF_CSRF_ENABLED': csrf_enabled,
    })


def test_csrf_enabled_by_default():
    app = make_app()
    assert app.config['WTF_CSRF_ENABLED'] is True


def test_browser_form_post_without_csrf_is_rejected():
    app = make_app()
    with app.test_client() as client:
        response = client.post('/auth/login', data={
            'username': 'student',
            'password': 'Password123!'
        })
    assert response.status_code == 400


def test_json_analyse_api_is_csrf_exempt_for_postman_testing():
    app = make_app()
    with app.test_client() as client:
        response = client.post('/api/analyse', json={})
    assert response.status_code == 400
    assert response.get_json()['message'] == 'email_text is required'


def test_flask_debug_mode_is_not_hardcoded_true():
    source = cybermailguard_app.__loader__.get_source(cybermailguard_app.__name__)
    assert 'app.run(debug=True)' not in source
    assert 'FLASK_DEBUG' in source


def test_2fa_rejects_invalid_token():
    secret = generate_secret()
    assert verify_token(secret, '000000') is False


def test_ml_training_rejects_csv_without_required_columns(tmp_path):
    bad_csv = tmp_path / 'bad_dataset.csv'
    bad_csv.write_text('message,type\nhello,safe\n', encoding='utf-8')

    result = train_model(str(bad_csv))

    assert result['error'] == 'CSV must have "text" and "label" columns'


def test_pdf_export_blocks_access_to_other_users_analysis(monkeypatch):
    app = make_app(csrf_enabled=False)
    analysis_id = ObjectId()
    owner_id = ObjectId()
    current_user_id = ObjectId()

    class FakeCollection:
        def find_one(self, query):
            assert query == {'_id': analysis_id}
            return {
                '_id': analysis_id,
                'user_id': owner_id,
                'email_subject': 'Private analysis',
                'risk_level': 'GREEN',
                'risk_score': 10,
                'reasons': [],
            }

    class FakeDb(dict):
        def __getitem__(self, name):
            return FakeCollection()

    monkeypatch.setattr(dashboard_routes, 'get_db', lambda: FakeDb())
    monkeypatch.setattr(dashboard_routes, 'current_user', SimpleNamespace(
        id=str(current_user_id),
        role='analyst',
    ))

    with app.test_request_context(f'/export/pdf/{analysis_id}'):
        response = dashboard_routes.export_pdf.__wrapped__(str(analysis_id))

    assert response.status_code == 302


def test_csv_export_requires_user_analysis(monkeypatch):
    app = make_app(csrf_enabled=False)
    current_user_id = ObjectId()

    class EmptyCursor(list):
        pass

    class FakeCollection:
        def find(self, query):
            assert query == {'user_id': current_user_id}
            return EmptyCursor()

    class FakeDb(dict):
        def __getitem__(self, name):
            return FakeCollection()

    monkeypatch.setattr(dashboard_routes, 'get_db', lambda: FakeDb())
    monkeypatch.setattr(dashboard_routes, 'current_user', SimpleNamespace(
        id=str(current_user_id),
        role='analyst',
    ))

    with app.test_request_context('/export/csv'):
        response = dashboard_routes.export_csv.__wrapped__()

    assert response.status_code == 302
    assert response.location.endswith('/analyse')


def test_csv_export_requires_admin_analysis(monkeypatch):
    app = make_app(csrf_enabled=False)

    class EmptyCursor(list):
        pass

    class FakeCollection:
        def find(self, query):
            assert query == {}
            return EmptyCursor()

    class FakeDb(dict):
        def __getitem__(self, name):
            return FakeCollection()

    monkeypatch.setattr(dashboard_routes, 'get_db', lambda: FakeDb())
    monkeypatch.setattr(dashboard_routes, 'current_user', SimpleNamespace(
        id=str(ObjectId()),
        role='admin',
    ))

    with app.test_request_context('/export/csv'):
        response = dashboard_routes.export_csv.__wrapped__()

    assert response.status_code == 302
    assert response.location.endswith('/analyse')


def test_docx_upload_text_can_be_extracted():
    doc = Document()
    doc.add_paragraph('Subject: Urgent account verification')
    doc.add_paragraph('Click http://secure-bank-verification-login.com')
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    raw = dashboard_routes._raw_text_from_upload(buffer, 'sample.docx')

    assert 'Urgent account verification' in raw
    assert 'http://secure-bank-verification-login.com' in raw
